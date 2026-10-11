import csv
import io

import pytest

from app.api.v1.settings import sanitize_spreadsheet_cell, write_csv_row, write_xlsx_row


@pytest.mark.parametrize(
    "value",
    [
        '=HYPERLINK("https://example.invalid")',
        "+SUM(1,2)",
        "-1+2",
        "@SUM(1,2)",
        " \t\r\n=1+1",
        "\ufeff=1+1",
        "\u200b=1+1",
        " \n\t@SUM(1,2)",
        "＝1＋1",
        "－1＋1",
        "＠SUM(1,2)",
    ],
)
def test_formula_like_text_is_made_literal(value):
    assert sanitize_spreadsheet_cell(value) == "'" + value


@pytest.mark.parametrize("value", ["Oats", "café 🍎", "a,b", 'say "hi"', "line1\nline2", ""])
def test_normal_unicode_and_csv_text_is_preserved(value):
    assert sanitize_spreadsheet_cell(value) == value


def test_negative_numbers_and_non_text_values_are_preserved():
    assert sanitize_spreadsheet_cell("-42") == "-42"
    assert sanitize_spreadsheet_cell("-12.5") == "-12.5"
    assert sanitize_spreadsheet_cell(0) == 0
    assert sanitize_spreadsheet_cell(None) is None


def test_row_writer_keeps_csv_escaping_and_sanitizes_each_text_cell():
    output = io.StringIO()
    writer = csv.writer(output, lineterminator="\r\n")
    write_csv_row(writer, ["Task", "=1+1", 'comma, quote " and newline\n'])

    parsed = next(csv.reader(io.StringIO(output.getvalue())))
    assert parsed == ["Task", "'=1+1", 'comma, quote " and newline\n']


def test_xlsx_writer_saves_formula_like_values_as_literal_strings():
    from openpyxl import Workbook, load_workbook

    workbook = Workbook()
    sheet = workbook.active
    write_xlsx_row(sheet, ["Food", "=1+1"])
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    cell = load_workbook(output, data_only=False).active["B1"]
    assert cell.value == "'=1+1"
    assert cell.data_type == "s"
